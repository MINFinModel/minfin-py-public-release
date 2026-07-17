"""Adaptive pure-input scenario labels from COVER!C10."""

from __future__ import annotations

import pandas as pd
import pytest

from MinFin.pure_input_blocks import (
    build_input_blocks_from_pure_input_file,
    pure_input_scenario_label,
    read_cover_scenario,
)


def _write_cover(path, *, b10: str, c10):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "COVER"
    ws["B10"] = b10
    ws["C10"] = c10
    wb.save(path)


def _write_investment_plan(path, scenarios: list[str], cover_scenario=None):
    from openpyxl import Workbook

    years = [2025, 2026]
    rows = []
    for scen in scenarios:
        rows.append(
            {
                "Variable": "Capital Cost",
                "Scenario": scen,
                "Technology": "TechA",
                "Unit": "Million USD",
                2025: 10.0,
                2026: 20.0,
            }
        )
        rows.append(
            {
                "Variable": "ActualGeneration",
                "Scenario": scen,
                "Technology": "TechA",
                "Unit": "PJ",
                2025: 1.0,
                2026: 2.0,
            }
        )
        rows.append(
            {
                "Variable": "OPEX",
                "Scenario": scen,
                "Technology": "TechA",
                "Unit": "Million USD",
                2025: 0.5,
                2026: 0.6,
            }
        )
        rows.append(
            {
                "Variable": "PotentialGeneration",
                "Scenario": scen,
                "Technology": "TechA",
                "Unit": "PJ",
                2025: 1.5,
                2026: 2.5,
            }
        )
    plan = pd.DataFrame(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "COVER"
    if cover_scenario is not None:
        ws["B10"] = "Scenario"
        ws["C10"] = cover_scenario
    else:
        ws["B10"] = "Notes"
        ws["C10"] = None
    wb.create_sheet("INVESTMENT PLAN")
    wb.save(path)
    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        plan.to_excel(writer, sheet_name="INVESTMENT PLAN", startrow=3, index=False)


def test_read_cover_scenario_requires_scenario_label(tmp_path):
    path = tmp_path / "cover.xlsx"
    _write_cover(path, b10="Notes", c10="Mitigation")
    assert read_cover_scenario(str(path)) is None

    _write_cover(path, b10="Scenario", c10="Mitigation")
    read_cover_scenario.cache_clear()
    assert read_cover_scenario(str(path)) == "Mitigation"


def test_pure_input_scenario_label_uses_cover_when_default_absent(tmp_path):
    path = tmp_path / "mitigation.xlsx"
    _write_investment_plan(path, scenarios=["Mitigation"], cover_scenario="Mitigation")
    read_cover_scenario.cache_clear()

    assert pure_input_scenario_label("net_zero", str(path)) == "Mitigation"
    assert pure_input_scenario_label("least_cost", str(path)) == "Mitigation"
    assert pure_input_scenario_label("Net Zero", str(path)) == "Mitigation"


def test_pure_input_scenario_label_keeps_defaults_when_present(tmp_path):
    path = tmp_path / "nz.xlsx"
    _write_investment_plan(
        path, scenarios=["Net Zero", "Least Cost"], cover_scenario="Net Zero"
    )
    read_cover_scenario.cache_clear()

    assert pure_input_scenario_label("net_zero", str(path)) == "Net Zero"
    assert pure_input_scenario_label("least_cost", str(path)) == "Least Cost"


def test_pure_input_scenario_label_uses_sole_plan_scenario_when_cover_blank(tmp_path):
    """Zambia-style: COVER!B10 is Scenario but C10 empty; plan only has IRP."""
    path = tmp_path / "irp.xlsx"
    _write_investment_plan(path, scenarios=["IRP"], cover_scenario=None)
    # Match Zambia cover layout: label present, value blank
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["COVER"]
    ws["B10"] = "Scenario"
    ws["C10"] = None
    wb.save(path)
    read_cover_scenario.cache_clear()

    assert read_cover_scenario(str(path)) is None
    assert pure_input_scenario_label("net_zero", str(path)) == "IRP"
    blocks = build_input_blocks_from_pure_input_file("net_zero", str(path))
    assert not blocks["elec_production"].empty
    assert "TechA" in blocks["elec_production"].columns
