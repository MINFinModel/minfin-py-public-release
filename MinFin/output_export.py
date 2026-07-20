"""Helpers for exporting notebook outputs to Excel workbooks."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from MinFin.excel_io import year_columns_from_index
from MinFin.offtaker_tariffs import _classify_segment, _get_tech_class
from MinFin.technology_sheet_io import PURE_INPUT_STATIC_FIELD_SOURCES


RAW_INPUT_COLUMNS = [
    "Scenario",
    "Variable",
    "technology",
    "year",
    "unit",
    "ResultValue",
]

FINANCING_SOURCES = ("Comm_Intl", "Comm_Dom", "Conc_IFI", "Conc_DPS")

PPA_INPUT_VARIABLES = {
    name
    for name, (sheet, _var) in PURE_INPUT_STATIC_FIELD_SOURCES.items()
    if sheet == "PPA REVENUE"
} | {"ppa_currency"}

PPA_COMPUTED_VARIABLES = {
    "total_ppa_revenue",
    "ppa_met_generation",
    "ppa_direct_offtaker_share",
}

WHOLESALE_COMPUTED_VARIABLES = {
    "whole_sale_generation",
    "total_wholesale_revnue",
    "sale_price",
}

NETWORK_SEGMENT_VARIABLES = {
    "power_purchased",
    "tariff",
    "capacity_purchased",
    "capacity_tariff",
    "power_purchase_cost",
}

INVESTMENT_ALLOCATION_PREFIXES = ("local_currency_", "foreign_currency_")

# Legacy notebook columns used spaces; canonical export names use underscores.
CANONICAL_VARIABLE_ALIASES: dict[str, str] = {
    "power purchased": "power_purchased",
    "capacity purchased": "capacity_purchased",
    "capacity tariff": "capacity_tariff",
}

INVESTMENT_BLOCK_VARIABLES = (
    "capital_cost",
    "elec_production",
    "potential_generation",
    "opex",
    # "ffe",  # Temporarily excluded: pure-input workbooks have no FFE series (zeros only).
    "co2_emission",
    "emission_savings",
    "carbon_price",
)

ECONOMY_DIM = "Economy"

# Weighted-average financing terms per technology, exported as scalar rows
# (one value per technology, no year). Maps export variable -> (summary column, unit).
TECHNOLOGY_FINANCING_SUMMARY_FIELDS: dict[str, tuple[str, str]] = {
    "weighted_grace_period": ("Grace Period (years)", "years"),
    "weighted_loan_term": ("Loan Term (years)", "years"),
    "weighted_interest_rate": ("Combined Interest Rate (%)", "%"),
    "weighted_rate_of_return_on_equity": ("Equity Return Rate (%)", "%"),
    "weighted_wacc": ("WACC (%)", "%"),
}

ECONOMY_METRIC_VARIABLES = (
    "financing_baseline",
    "existing_financing_requirement",
    "financing_requirement_share_of_gdp",
    "funding_availability_share_of_gdp",
)

# Computed outputs inherit units from related workbook variables when possible.
COMPUTED_VARIABLE_UNIT_SOURCES: dict[str, str] = {
    "investment_need": "capital_cost",
    "total_grant_amount": "total_grant_amount",
    "total_generation": "ppa_contracted_generation",
    "ppa_met_generation": "ppa_contracted_generation",
    "whole_sale_generation": "ppa_contracted_generation",
    "ppa_direct_offtaker_share": "ppa_standard_offtaker_share",
    "total_ppa_revenue": "total_grant_amount",
    "total_wholesale_revnue": "total_grant_amount",
    "redispatch_compensation": "total_grant_amount",
    "opex": "opex",
    "power_purchased": "ppa_contracted_generation",
    "capacity_purchased": "ppa_contracted_capacity",
    "power_purchase_cost": "total_grant_amount",
    "corporate_tax_expenses": "total_grant_amount",
    "interest_cost": "total_grant_amount",
    "cashflow": "total_grant_amount",
    "Financing Requirement": "total_grant_amount",
    "Existing Financing Requirement": "total_grant_amount",
    "total_financing_requirement": "total_grant_amount",
    "financing_baseline": "total_grant_amount",
    "existing_financing_requirement": "total_grant_amount",
}

# Share-of-GDP metrics are unitless ratios; they never inherit a money unit.
GDP_SHARE_VARIABLES = (
    "financing_requirement_share_of_gdp",
    "funding_availability_share_of_gdp",
)

DEFAULT_VARIABLE_UNITS: dict[str, str] = {
    "investment_need": "Mn USD",
    "total_generation": "GWh/Year",
    "total_grant_amount": "Mn USD",
    "total_ppa_revenue": "Mn USD",
    "ppa_met_generation": "GWh/Year",
    "whole_sale_generation": "GWh/Year",
    "total_wholesale_revnue": "Mn USD",
    "sale_price": "USD/kWh",
    "redispatch_compensation": "Mn USD",
    "opex": "Mn USD",
    "power_purchased": "GWh/Year",
    "tariff": "USD/kWh",
    "capacity_purchased": "MW",
    "capacity_tariff": "Mn USD/MW",
    "power_purchase_cost": "Mn USD",
    "corporate_tax_expenses": "Mn USD",
    "interest_cost": "Mn USD",
    "cashflow": "Mn USD",
    "Financing Requirement": "Mn USD",
    "Existing Financing Requirement": "Mn USD",
    "total_financing_requirement": "Mn USD",
    "financing_baseline": "Mn USD",
    "existing_financing_requirement": "Mn USD",
    "financing_requirement_share_of_gdp": "share of GDP",
    "funding_availability_share_of_gdp": "share of GDP",
}


def _normalize_years(years: Iterable | None, values: list | tuple | pd.Series) -> list:
    if years is None:
        return list(range(len(values)))
    years_list = list(years)
    if len(years_list) >= len(values):
        return years_list[: len(values)]
    return years_list + list(range(len(years_list), len(values)))


def _values_are_meaningful(values: list | tuple | pd.Series) -> bool:
    """True when at least one value is non-null and non-zero."""
    series = pd.to_numeric(pd.Series(list(values)), errors="coerce")
    if series.notna().sum() == 0:
        return False
    return bool((series.fillna(0) != 0).any())


def _is_zero_currency_placeholder(variable: str, values: list | tuple | pd.Series) -> bool:
    if variable != "ppa_currency":
        return False
    return not _values_are_meaningful(values)


def _is_wholesale_variable(variable: str) -> bool:
    return "_share_" in variable or "_sale_price_" in variable


def _is_investment_allocation_variable(variable: str) -> bool:
    return variable.startswith(INVESTMENT_ALLOCATION_PREFIXES)


def _investment_allocation_unit(
    variable: str,
    *,
    display_currency: str | None = None,
    local_currency_code: str | None = None,
    foreign_currency_code: str | None = None,
) -> str:
    """Unit label for ``local_currency_*`` / ``foreign_currency_*`` columns.

    Prefer MACROECONOMIC **Display Currency** for both prefixes. When unset,
    fall back to local/foreign codes, then ``USD``.
    """
    if variable.startswith("local_currency_"):
        code = display_currency or local_currency_code or "USD"
        return f"Mn {code}"
    if variable.startswith("foreign_currency_"):
        code = display_currency or foreign_currency_code or "USD"
        return f"Mn {code}"
    return ""


def _allocation_native_currency(
    variable: str,
    *,
    local_currency_code: str | None = None,
    foreign_currency_code: str | None = None,
) -> str | None:
    """Currency the allocation column is denominated in before export conversion."""
    if variable.startswith("local_currency_"):
        return local_currency_code
    if variable.startswith("foreign_currency_"):
        return foreign_currency_code
    return None


def _convert_allocation_values_to_display(
    values: list,
    years: Iterable,
    *,
    source_currency: str | None,
    display_currency: str | None,
    exchange_rates: Optional[pd.DataFrame],
) -> list:
    """Convert allocation amounts from *source_currency* to *display_currency* per year.

    Rates are MACROECONOMIC-style (same units for all currencies, e.g. LCU per USD).
    Conversion: ``value * display_rate / source_rate``. No-op when currencies match
    or rates are unavailable.
    """
    from MinFin.fx import normalize_currency_code

    if not values:
        return values
    src_raw = (source_currency or "").strip()
    dst_raw = (display_currency or "").strip()
    if not src_raw or not dst_raw or src_raw == dst_raw:
        return values
    if exchange_rates is None or not isinstance(exchange_rates, pd.DataFrame) or exchange_rates.empty:
        return values

    # Normalize year index so int/float/str year labels all resolve.
    rates = exchange_rates.copy()
    try:
        rates.index = pd.to_numeric(rates.index, errors="coerce")
        rates = rates[rates.index.notna()]
        rates.index = rates.index.astype(int)
    except Exception:
        rates = exchange_rates

    available = [str(c).strip() for c in rates.columns]
    rates.columns = available
    src = normalize_currency_code(src_raw, available)
    dst = normalize_currency_code(dst_raw, available)
    if src not in rates.columns or dst not in rates.columns:
        return values
    if src == dst:
        return values

    year_list = list(years) if years is not None else []
    if len(year_list) < len(values):
        # Prefer caller-supplied years; if too short/empty, they are unusable for FX.
        # Caller should pass tech_df.index — fall back to positional only as last resort.
        year_list = list(year_list) + [None] * (len(values) - len(year_list))

    out: list = []
    for i, value in enumerate(values):
        if value is None or (isinstance(value, float) and pd.isna(value)):
            out.append(value)
            continue
        try:
            year = int(float(year_list[i]))
        except (IndexError, TypeError, ValueError):
            out.append(value)
            continue
        if year not in rates.index:
            out.append(value)
            continue
        src_rate = rates.loc[year, src]
        dst_rate = rates.loc[year, dst]
        if isinstance(src_rate, pd.Series):
            src_rate = src_rate.iloc[0]
        if isinstance(dst_rate, pd.Series):
            dst_rate = dst_rate.iloc[0]
        if pd.isna(src_rate) or pd.isna(dst_rate) or float(src_rate) == 0.0:
            out.append(value)
            continue
        out.append(float(value) * float(dst_rate) / float(src_rate))
    return out


def _canonical_variable_name(variable: str) -> str:
    return CANONICAL_VARIABLE_ALIASES.get(str(variable).strip(), str(variable).strip())


def _iter_canonical_tech_columns(tech_df: pd.DataFrame) -> list[tuple[str, str]]:
    """Return (canonical_name, source_column) pairs; later columns win on alias clashes."""
    mapping: dict[str, str] = {}
    for column in tech_df.columns:
        mapping[_canonical_variable_name(column)] = column
    return sorted((canonical, source) for canonical, source in mapping.items())


def _parent_tech_block_values(
    block_df: pd.DataFrame,
    parent_tech: str,
    df_technologies: Optional[pd.DataFrame],
) -> Optional[list]:
    if block_df is None or block_df.empty or df_technologies is None:
        return None
    if "Technology" not in df_technologies.columns or "Name" not in df_technologies.columns:
        return None

    codes = (
        df_technologies.loc[df_technologies["Technology"] == parent_tech, "Name"]
        .dropna()
        .astype(str)
        .str.strip()
        .tolist()
    )
    if not codes:
        return None

    data = block_df.copy()
    if "Year" in data.columns:
        data = data.set_index("Year")
    data = data.loc[data.index != "Total"] if "Total" in data.index else data

    cols = [code for code in codes if code in data.columns]
    if not cols:
        return None

    series = data[cols].apply(pd.to_numeric, errors="coerce").fillna(0).sum(axis=1)
    return series.tolist()


def _is_network_segment(tech_name: str, df_technologies: Optional[pd.DataFrame]) -> bool:
    if df_technologies is None:
        return False
    segment = _classify_segment(_get_tech_class(tech_name, df_technologies))
    return segment >= 2


def _tech_has_ppa_inputs(tech_name: str, all_tech_data: Optional[dict]) -> bool:
    if not all_tech_data or tech_name not in all_tech_data:
        return False
    tech_fields = all_tech_data[tech_name]
    if not isinstance(tech_fields, dict):
        return False
    for variable in PPA_INPUT_VARIABLES:
        if variable == "ppa_currency":
            continue
        payload = tech_fields.get(variable)
        if isinstance(payload, dict) and _values_are_meaningful(payload.get("values", [])):
            return True
    return False


def _tech_has_wholesale_inputs(tech_name: str, all_tech_data: Optional[dict]) -> bool:
    if not all_tech_data or tech_name not in all_tech_data:
        return False
    tech_fields = all_tech_data[tech_name]
    if not isinstance(tech_fields, dict):
        return False
    for variable, payload in tech_fields.items():
        if not _is_wholesale_variable(variable):
            continue
        if isinstance(payload, dict) and _values_are_meaningful(payload.get("values", [])):
            return True
    return False


def _should_export_variable(
    variable: str,
    values: list | tuple | pd.Series,
    *,
    tech_name: str,
    all_tech_data: Optional[dict] = None,
    df_technologies: Optional[pd.DataFrame] = None,
    from_tech_dataframe: bool = False,
) -> bool:
    """Return whether *variable* applies to *tech_name*.

    Applicability is based on technology type (PPA / wholesale / network), not on
    whether the series is all-zero. Zero-valued but applicable rows are still exported.
    """
    if _is_zero_currency_placeholder(variable, values):
        return False

    canonical = _canonical_variable_name(variable)
    if canonical in PPA_INPUT_VARIABLES or canonical in PPA_COMPUTED_VARIABLES:
        if from_tech_dataframe:
            return _tech_has_ppa_inputs(tech_name, all_tech_data) or _values_are_meaningful(values)
        return _tech_has_ppa_inputs(tech_name, all_tech_data)
    if _is_wholesale_variable(canonical) or canonical in WHOLESALE_COMPUTED_VARIABLES:
        if from_tech_dataframe:
            return True
        return _tech_has_wholesale_inputs(tech_name, all_tech_data)
    if canonical in NETWORK_SEGMENT_VARIABLES:
        return _is_network_segment(tech_name, df_technologies)
    return True


def _lookup_workbook_unit(variable: str, variable_units: Optional[dict[str, str]]) -> str:
    if not variable_units:
        return ""
    return variable_units.get(variable, "")


def _variable_unit(
    variable: str,
    *,
    all_tech_data: Optional[dict] = None,
    tech_name: Optional[str] = None,
    variable_units: Optional[dict[str, str]] = None,
    display_currency: str | None = None,
    local_currency_code: str | None = None,
    foreign_currency_code: str | None = None,
) -> str:
    if _is_investment_allocation_variable(variable):
        return _investment_allocation_unit(
            variable,
            display_currency=display_currency,
            local_currency_code=local_currency_code,
            foreign_currency_code=foreign_currency_code,
        )

    if all_tech_data and tech_name and tech_name in all_tech_data:
        payload = all_tech_data[tech_name].get(variable)
        if isinstance(payload, dict):
            unit = payload.get("unit", "")
            if unit:
                return unit

    workbook_unit = _lookup_workbook_unit(variable, variable_units)
    if workbook_unit:
        return workbook_unit

    source = COMPUTED_VARIABLE_UNIT_SOURCES.get(variable)
    if source:
        if all_tech_data and tech_name and tech_name in all_tech_data:
            payload = all_tech_data[tech_name].get(source)
            if isinstance(payload, dict):
                unit = payload.get("unit", "")
                if unit:
                    return unit
        inherited = _lookup_workbook_unit(source, variable_units)
        if inherited:
            return inherited

    if variable.startswith("Loans (") or variable.startswith("Equity ("):
        money_unit = _lookup_workbook_unit("total_grant_amount", variable_units)
        return money_unit or DEFAULT_VARIABLE_UNITS["Financing Requirement"]

    return DEFAULT_VARIABLE_UNITS.get(variable, "")


def _repayment_schedule_year_columns(
    repayment_schedule: pd.DataFrame,
    *,
    since_year: int | None = None,
) -> list[int]:
    return year_columns_from_index(repayment_schedule.columns, since_year=since_year)


def compute_financing_baseline_by_year(repayment_schedule: pd.DataFrame) -> pd.Series:
    """Annual repayment totals from the historic financing-baseline portfolio (all years)."""
    year_cols = _repayment_schedule_year_columns(repayment_schedule)
    if not year_cols:
        return pd.Series(dtype=float, name="financing_baseline")
    totals = repayment_schedule[year_cols].sum()
    totals.index = totals.index.astype(int)
    return totals.sort_index().rename("financing_baseline")


def compute_existing_financing_requirement_by_year(
    repayment_schedule: pd.DataFrame,
    years: Iterable | None = None,
) -> pd.Series:
    """Annual repayment totals for model projection years (existing debt service)."""
    if years is not None:
        year_set = {int(y) for y in years}
        year_cols = [y for y in _repayment_schedule_year_columns(repayment_schedule) if y in year_set]
    else:
        year_cols = _repayment_schedule_year_columns(repayment_schedule)
    if not year_cols:
        return pd.Series(dtype=float, name="existing_financing_requirement")
    totals = repayment_schedule[year_cols].sum()
    totals.index = totals.index.astype(int)
    return totals.sort_index().rename("existing_financing_requirement")


def compute_existing_financing_requirement_by_technology(
    repayment_schedule: pd.DataFrame,
    technology: pd.Series | Iterable | None,
    years: Iterable | None = None,
    technology_map: dict[str, str] | None = None,
) -> pd.DataFrame:
    """Annual existing-debt service split by technology.

    The historic financing-baseline portfolio (``repayment_schedule``) carries no
    ``Technology`` column, so the caller passes the aligned ``historical['Technology']``
    labels. ``technology_map`` (e.g. from ``load_technology_alias_map``) optionally renames
    each label to the parent model technology so existing financing lines up with the rest of
    the outputs. Returns a DataFrame indexed by technology with one column per projection year.
    """
    if repayment_schedule is None or repayment_schedule.empty or technology is None:
        return pd.DataFrame()

    year_cols = _repayment_schedule_year_columns(repayment_schedule)
    if years is not None:
        year_set = {int(y) for y in years}
        year_cols = [y for y in year_cols if y in year_set]
    if not year_cols:
        return pd.DataFrame()

    tech = pd.Series(technology)
    if not tech.index.equals(repayment_schedule.index):
        if len(tech) == len(repayment_schedule):
            tech = pd.Series(tech.to_numpy(), index=repayment_schedule.index)
        else:
            tech = tech.reindex(repayment_schedule.index)

    if technology_map:
        tech = tech.map(lambda v: technology_map.get(str(v).strip(), v))

    grouped = repayment_schedule[year_cols].groupby(tech).sum()
    grouped.columns = [int(c) for c in grouped.columns]
    grouped = grouped.sort_index(axis=1)
    grouped.index.name = "Technology"
    return grouped


def _append_economy_metric_records(
    records: list[dict],
    *,
    variable: str,
    series: pd.Series,
    scenario: str,
    unit: str,
    include_unit_dim: bool,
) -> None:
    if series is None or series.empty:
        return
    _append_records(
        records,
        variable=variable,
        tech_name=ECONOMY_DIM,
        values=series.tolist(),
        years=series.index.tolist(),
        scenario=scenario,
        unit=unit,
        include_unit_dim=include_unit_dim,
    )


def technology_financing_summary_from_weighted_averages(
    weighted_averages: pd.DataFrame,
) -> pd.DataFrame:
    """Build export summary from NEW INFRASTRUCTURE ``weighted_averages`` (Technology Disag top table).

    ``weighted_averages`` uses fractional rates (0.07 = 7%); export columns use percent (7.0).
    """
    if weighted_averages is None or getattr(weighted_averages, "empty", True):
        return pd.DataFrame()

    wa = weighted_averages

    def _col(category: str, parameter: str) -> pd.Series:
        key = (category, parameter)
        if key in wa.columns:
            return wa[key]
        return pd.Series(0.0, index=wa.index)

    debt_ir = _col("Debt", "Interest Rate")
    equity_ror = _col("Equity", "Rate of Return")
    debt_share = _col("Financing Shares", "Debt Share")
    equity_share = _col("Financing Shares", "Equity Share")

    summary = pd.DataFrame(index=wa.index)
    summary.index.name = "Technology"
    summary["Grace Period (years)"] = _col("Debt", "Grace Period")
    summary["Loan Term (years)"] = _col("Debt", "Loan Term")
    summary["Equity Return Rate (%)"] = equity_ror * 100
    summary["Combined Interest Rate (%)"] = (
        debt_ir * debt_share + equity_ror * equity_share
    ) * 100
    if ("Weighted Cost of Capital", "WACC") in wa.columns:
        summary["WACC (%)"] = wa[("Weighted Cost of Capital", "WACC")] * 100
    else:
        summary["WACC (%)"] = summary["Combined Interest Rate (%)"]
    return summary


def _append_technology_financing_summary_records(
    records: list[dict],
    *,
    summary: pd.DataFrame,
    scenario: str,
    include_unit_dim: bool,
) -> None:
    """Emit weighted-average financing terms (term, grace, rates, RoE, WACC) per technology.

    Each metric is a single weighted average per technology, so ``year`` is left blank.
    """
    if summary is None or getattr(summary, "empty", True):
        return

    for variable, (column, unit) in TECHNOLOGY_FINANCING_SUMMARY_FIELDS.items():
        if column not in summary.columns:
            continue
        for tech_name, value in summary[column].items():
            if pd.isna(value):
                value = None
            records.append(
                {
                    "Scenario": scenario,
                    "Variable": variable,
                    "technology": tech_name,
                    "year": None,
                    "unit": unit if include_unit_dim else None,
                    "ResultValue": value,
                }
            )


def _append_records(
    records: list[dict],
    *,
    variable: str,
    tech_name: str,
    values: list | tuple | pd.Series,
    years: Iterable | None,
    scenario: str,
    unit: str,
    include_unit_dim: bool,
) -> None:
    row_years = _normalize_years(years, values)
    for year, value in zip(row_years, values):
        if pd.isna(value):
            value = None
        records.append(
            {
                "Scenario": scenario,
                "Variable": variable,
                "technology": tech_name,
                "year": year,
                "unit": unit if include_unit_dim else None,
                "ResultValue": value,
            }
        )


def build_technology_parameter_raw_table(
    all_tech_data: dict,
    years: Iterable | None = None,
    scenario: str = "Net Zero",
    include_unit_dim: bool = True,
    variable_units: Optional[dict[str, str]] = None,
) -> pd.DataFrame:
    """Flatten ``all_tech_data`` into a raw-input-style long table.

    Output columns use human-readable names instead of the OSeMOSYS template's
    generic ``Dim`` columns:

    - ``Variable`` stores the MinFin technology parameter name.
    - ``technology`` stores the technology name.
    - ``year`` stores the year.
    - ``unit`` stores the unit when ``include_unit_dim`` is enabled.
    """
    records: list[dict] = []

    for tech_name, tech_fields in (all_tech_data or {}).items():
        if not isinstance(tech_fields, dict):
            continue
        for variable, payload in tech_fields.items():
            if not isinstance(payload, dict):
                continue

            values = list(payload.get("values", []))
            if not values or _is_zero_currency_placeholder(variable, values):
                continue
            if not _should_export_variable(
                _canonical_variable_name(variable),
                values,
                tech_name=tech_name,
                all_tech_data=all_tech_data,
            ):
                continue

            unit = payload.get("unit", "") or _lookup_workbook_unit(variable, variable_units)
            _append_records(
                records,
                variable=variable,
                tech_name=tech_name,
                values=values,
                years=years,
                scenario=scenario,
                unit=unit,
                include_unit_dim=include_unit_dim,
            )

    df = pd.DataFrame(records, columns=RAW_INPUT_COLUMNS)
    if df.empty:
        return df

    return df.sort_values(["Variable", "technology", "year"], kind="stable").reset_index(drop=True)


def build_technology_output_raw_table(
    tech_dataframes: dict,
    *,
    all_tech_data: Optional[dict] = None,
    financing_requirement_by_tech: Optional[dict] = None,
    investment_blocks: Optional[dict[str, pd.DataFrame]] = None,
    df_technologies: Optional[pd.DataFrame] = None,
    variable_units: Optional[dict[str, str]] = None,
    years: Iterable | None = None,
    repayment_schedule: Optional[pd.DataFrame] = None,
    economy_metrics: Optional[dict[str, pd.Series]] = None,
    technology_financing_summary: Optional[pd.DataFrame] = None,
    existing_financing_by_technology: Optional[pd.DataFrame] = None,
    scenario: str = "Net Zero",
    include_unit_dim: bool = True,
    display_currency: str | None = None,
    local_currency_code: str | None = None,
    foreign_currency_code: str | None = None,
    exchange_rates: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """Flatten per-technology outputs (inputs + computed) into a long table.

    Only variables that are meaningful for a given technology are exported.
    For example, PPA parameters are omitted for technologies without PPA inputs,
    and network purchase-cost fields are limited to Transmission/Distribution/Exports.

    ``local_currency_*`` / ``foreign_currency_*`` amounts are converted to
    *display_currency* (when set) using *exchange_rates* so ResultValue matches
    the unit label.
    """
    records: list[dict] = []
    exported: set[tuple[str, str]] = set()

    for tech_name, tech_df in (tech_dataframes or {}).items():
        if not isinstance(tech_df, pd.DataFrame) or tech_df.empty:
            continue

        for variable, source_column in _iter_canonical_tech_columns(tech_df):
            column_data = tech_df[source_column]
            if isinstance(column_data, pd.DataFrame):
                column_data = column_data.iloc[:, -1]
            values = pd.to_numeric(column_data, errors="coerce").tolist()
            if not _should_export_variable(
                variable,
                values,
                tech_name=tech_name,
                all_tech_data=all_tech_data,
                df_technologies=df_technologies,
                from_tech_dataframe=True,
            ):
                continue

            row_years = tech_df.index.tolist()
            unit = _variable_unit(
                variable,
                all_tech_data=all_tech_data,
                tech_name=tech_name,
                variable_units=variable_units,
                display_currency=display_currency,
                local_currency_code=local_currency_code,
                foreign_currency_code=foreign_currency_code,
            )
            if _is_investment_allocation_variable(variable):
                unit_code = (
                    display_currency
                    or (
                        local_currency_code
                        if variable.startswith("local_currency_")
                        else foreign_currency_code
                    )
                    or "USD"
                )
                values = _convert_allocation_values_to_display(
                    values,
                    row_years,
                    source_currency=_allocation_native_currency(
                        variable,
                        local_currency_code=local_currency_code,
                        foreign_currency_code=foreign_currency_code,
                    ),
                    display_currency=unit_code,
                    exchange_rates=exchange_rates,
                )
            _append_records(
                records,
                variable=variable,
                tech_name=tech_name,
                values=values,
                years=row_years,
                scenario=scenario,
                unit=unit,
                include_unit_dim=include_unit_dim,
            )
            exported.add((tech_name, variable))

    if all_tech_data:
        for tech_name, tech_fields in all_tech_data.items():
            if not isinstance(tech_fields, dict):
                continue
            for variable, payload in tech_fields.items():
                canonical = _canonical_variable_name(variable)
                if (tech_name, canonical) in exported:
                    continue
                if not isinstance(payload, dict):
                    continue
                values = list(payload.get("values", []))
                if not values or _is_zero_currency_placeholder(canonical, values):
                    continue
                if not _should_export_variable(
                    canonical,
                    values,
                    tech_name=tech_name,
                    all_tech_data=all_tech_data,
                    df_technologies=df_technologies,
                ):
                    continue
                _append_records(
                    records,
                    variable=canonical,
                    tech_name=tech_name,
                    values=values,
                    years=years,
                    scenario=scenario,
                    unit=payload.get("unit", "") or _lookup_workbook_unit(canonical, variable_units),
                    include_unit_dim=include_unit_dim,
                )
                exported.add((tech_name, canonical))

    for block_name, block_df in (investment_blocks or {}).items():
        if block_name not in INVESTMENT_BLOCK_VARIABLES:
            continue
        for tech_name in (tech_dataframes or {}):
            if (tech_name, block_name) in exported:
                continue
            values = _parent_tech_block_values(block_df, tech_name, df_technologies)
            if values is None:
                continue
            _append_records(
                records,
                variable=block_name,
                tech_name=tech_name,
                values=values,
                years=years,
                scenario=scenario,
                unit=_lookup_workbook_unit(block_name, variable_units),
                include_unit_dim=include_unit_dim,
            )
            exported.add((tech_name, block_name))

    for tech_name, repayment_df in (financing_requirement_by_tech or {}).items():
        if not isinstance(repayment_df, pd.DataFrame) or repayment_df.empty:
            continue
        for variable in repayment_df.index:
            row = repayment_df.loc[variable]
            row_years = [int(y) for y in row.index]
            values = row.tolist()
            _append_records(
                records,
                variable=str(variable),
                tech_name=tech_name,
                values=values,
                years=row_years,
                scenario=scenario,
                unit=_variable_unit(
                    str(variable),
                    all_tech_data=all_tech_data,
                    tech_name=tech_name,
                    variable_units=variable_units,
                    display_currency=display_currency,
                    local_currency_code=local_currency_code,
                    foreign_currency_code=foreign_currency_code,
                ),
                include_unit_dim=include_unit_dim,
            )

    has_existing_by_tech = (
        existing_financing_by_technology is not None
        and not existing_financing_by_technology.empty
    )

    economy_series = dict(economy_metrics or {})
    if repayment_schedule is not None and not repayment_schedule.empty:
        economy_series.setdefault(
            "financing_baseline",
            compute_financing_baseline_by_year(repayment_schedule),
        )
        # Existing financing requirement is reported per technology (see the notebook's
        # ``financing_requirement_by_tech`` "Existing Financing Requirement" rows and the
        # ``existing_financing_by_technology`` argument), so it is no longer emitted at the
        # economy level here.

    for variable in ECONOMY_METRIC_VARIABLES:
        series = economy_series.get(variable)
        if series is None or getattr(series, "empty", True):
            continue
        _append_economy_metric_records(
            records,
            variable=variable,
            series=series,
            scenario=scenario,
            unit=_variable_unit(variable, variable_units=variable_units),
            include_unit_dim=include_unit_dim,
        )

    if has_existing_by_tech:
        _existing_unit = _variable_unit(
            "existing_financing_requirement", variable_units=variable_units
        )
        for tech_name, tech_row in existing_financing_by_technology.iterrows():
            _append_records(
                records,
                variable="existing_financing_requirement",
                tech_name=str(tech_name),
                values=tech_row.tolist(),
                years=existing_financing_by_technology.columns.tolist(),
                scenario=scenario,
                unit=_existing_unit,
                include_unit_dim=include_unit_dim,
            )

    _append_technology_financing_summary_records(
        records,
        summary=technology_financing_summary,
        scenario=scenario,
        include_unit_dim=include_unit_dim,
    )

    df = pd.DataFrame(records, columns=RAW_INPUT_COLUMNS)
    if df.empty:
        return df

    return df.sort_values(["technology", "Variable", "year"], kind="stable").reset_index(drop=True)


def _format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, min(ws.max_column, 20) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for cell in ws[col_letter][: min(ws.max_row, 200)]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 35)

    wb.save(path)


def export_technology_parameter_raw_workbook(
    all_tech_data: dict,
    output_path: str | Path,
    years: Iterable | None = None,
    scenario: str = "Net Zero",
    sheet_name: str = "0.1 Raw data",
    variable_units: Optional[dict[str, str]] = None,
) -> Path:
    """Write a raw-input-style technology parameter workbook."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    raw_table = build_technology_parameter_raw_table(
        all_tech_data=all_tech_data,
        years=years,
        scenario=scenario,
        variable_units=variable_units,
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw_table.to_excel(writer, sheet_name=sheet_name, index=False)

    _format_workbook(output_path)
    return output_path


def export_technology_output_workbook(
    tech_dataframes: dict,
    output_path: str | Path,
    *,
    all_tech_data: Optional[dict] = None,
    financing_requirement_by_tech: Optional[dict] = None,
    investment_blocks: Optional[dict[str, pd.DataFrame]] = None,
    df_technologies: Optional[pd.DataFrame] = None,
    variable_units: Optional[dict[str, str]] = None,
    years: Iterable | None = None,
    repayment_schedule: Optional[pd.DataFrame] = None,
    economy_metrics: Optional[dict[str, pd.Series]] = None,
    technology_financing_summary: Optional[pd.DataFrame] = None,
    existing_financing_by_technology: Optional[pd.DataFrame] = None,
    scenario: str = "Net Zero",
    sheet_name: str = "0.1 Raw data",
    display_currency: str | None = None,
    local_currency_code: str | None = None,
    foreign_currency_code: str | None = None,
    exchange_rates: Optional[pd.DataFrame] = None,
) -> Path:
    """Write the full per-technology output table (inputs + computed variables)."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    raw_table = build_technology_output_raw_table(
        tech_dataframes=tech_dataframes,
        all_tech_data=all_tech_data,
        financing_requirement_by_tech=financing_requirement_by_tech,
        investment_blocks=investment_blocks,
        df_technologies=df_technologies,
        variable_units=variable_units,
        years=years,
        repayment_schedule=repayment_schedule,
        economy_metrics=economy_metrics,
        technology_financing_summary=technology_financing_summary,
        existing_financing_by_technology=existing_financing_by_technology,
        scenario=scenario,
        display_currency=display_currency,
        local_currency_code=local_currency_code,
        foreign_currency_code=foreign_currency_code,
        exchange_rates=exchange_rates,
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw_table.to_excel(writer, sheet_name=sheet_name, index=False)

    _format_workbook(output_path)
    return output_path
