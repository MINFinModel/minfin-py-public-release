"""Build infrastructure input blocks from the pure-input INVESTMENT PLAN sheet."""

from __future__ import annotations

import warnings
from functools import lru_cache
from typing import Optional

import numpy as np
import pandas as pd

from MinFin.excel_io import read_investment_plan_long, year_columns_from_dataframe


@lru_cache(maxsize=8)
def read_cover_scenario(file_path: str) -> Optional[str]:
    """Return the active scenario name from COVER!C10 when B10 is ``Scenario``.

    Pure-input country files (e.g. Türkiye) put the INVESTMENT PLAN / revenue
    Scenario label in COVER cell C10. Older Kenya-style covers use B10 for Notes
    and leave C10 blank — those return ``None`` so callers fall back to the
    hardcoded ``Net Zero`` / ``Least Cost`` labels.
    """
    try:
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception:
        return None
    try:
        if "COVER" not in wb.sheetnames:
            return None
        ws = wb["COVER"]
        if str(ws["B10"].value or "").strip().lower() != "scenario":
            return None
        value = ws["C10"].value
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return None
        text = str(value).strip()
        if not text or text.lower() in {"nan", "none"}:
            return None
        return text
    finally:
        wb.close()


def pure_input_scenario_label(scenario: str, file_path: Optional[str] = None) -> str:
    """Map internal scenario keys to workbook Scenario labels.

    Resolution order when *file_path* is set:

    1. Hardcoded default (``Net Zero`` / ``Least Cost``) if that label exists in
       INVESTMENT PLAN (keeps multi-scenario Kenya-style loads correct).
    2. COVER!C10 when B10 is ``Scenario`` and the cell is non-empty.
    3. The sole distinct INVESTMENT PLAN scenario (e.g. ``IRP``, ``Mitigation``)
       when COVER is blank and the defaults are absent.
    4. The hardcoded default.

    For **exports**, prefer :func:`read_cover_scenario` when C10 is set so the
    Scenario column matches the COVER label even if the plan still uses defaults.
    """
    defaults = {"net_zero": "Net Zero", "least_cost": "Least Cost"}
    default = defaults.get(scenario, scenario)
    if not file_path:
        return default

    try:
        found = {
            str(s).strip()
            for s in read_investment_plan_long(file_path)["Scenario"].dropna().unique()
            if str(s).strip() and str(s).strip().lower() not in {"nan", "none"}
        }
    except Exception:
        found = set()

    if default in found:
        return default

    cover = read_cover_scenario(file_path)
    if cover:
        return cover
    if len(found) == 1:
        return next(iter(found))
    return default


def export_scenario_label(file_path: Optional[str] = None) -> str:
    """Scenario name written to output tables: COVER!C10 when set, else adaptive plan label."""
    if file_path:
        cover = read_cover_scenario(file_path)
        if cover:
            return cover
    return pure_input_scenario_label("net_zero", file_path)


def pivot_investment_block(
    sub: pd.DataFrame, year_cols, n_rows: int = 56, block_name: str = ""
) -> pd.DataFrame:
    """
    Build one wide block: Year + one column per technology, plus Total row.
    """
    rows = []
    for _, row in sub.iterrows():
        tech = row.get("Technology")
        if pd.isna(tech) or str(tech).strip() == "":
            continue
        tech = str(tech).strip()
        for yc in year_cols:
            y = int(float(yc))
            v = row[yc]
            if pd.isna(v):
                v = 0
            rows.append({"Year": y, "Technology": tech, "value": v})
    if not rows:
        return pd.DataFrame()
    long = pd.DataFrame(rows)
    wide = long.pivot_table(index="Year", columns="Technology", values="value", aggfunc="sum")
    wide = wide.sort_index().reset_index()
    if len(wide) > n_rows:
        wide = wide.iloc[:n_rows]
    val_cols = [c for c in wide.columns if c != "Year"]
    total_row = pd.DataFrame(wide[val_cols].sum()).T
    total_row.insert(0, "Year", "Total")
    return pd.concat([wide, total_row], ignore_index=True)


def pure_input_ffe_zero_block(capital_block: pd.DataFrame) -> pd.DataFrame:
    """Zero block matching *capital_block* shape when FFE is absent in pure input."""
    if capital_block is None or capital_block.empty:
        return pd.DataFrame()
    df = capital_block.copy()
    for c in df.columns:
        if c != "Year" and c != "Total":
            df[c] = 0.0
    return df


def build_input_blocks_from_pure_input_file(scenario: str, file_path: str) -> dict:
    df = read_investment_plan_long(file_path)
    year_cols = year_columns_from_dataframe(df)
    if not year_cols:
        raise ValueError("INVESTMENT PLAN: no year columns found (expected 2025, 2026, …)")

    label = pure_input_scenario_label(scenario, file_path)
    if label not in set(df["Scenario"].dropna().astype(str).str.strip()):
        warnings.warn(
            f"INVESTMENT PLAN has no rows for scenario {label!r}. "
            f"Found: {list(df['Scenario'].dropna().unique())}. "
            f"Empty blocks may result.",
            stacklevel=2,
        )

    n_rows = 56

    def _block(variable_name: str) -> pd.DataFrame:
        sub = df[
            (df["Variable"] == variable_name)
            & (df["Scenario"].astype(str).str.strip() == label)
        ]
        return pivot_investment_block(sub, year_cols, n_rows=n_rows, block_name=variable_name)

    capital = _block("Capital Cost")
    elec = _block("ActualGeneration")
    opex = _block("OPEX")
    potential = _block("PotentialGeneration")
    ffe = pure_input_ffe_zero_block(capital)

    return {
        "capital_cost": capital,
        "ffe": ffe,
        "elec_production": elec,
        "opex": opex,
        "potential_generation": potential,
    }


def pure_input_other_input_series(
    file_path: str, scenario: str, year_cols, n_years: int = 46
) -> dict:
    label = pure_input_scenario_label(scenario, file_path)
    df = read_investment_plan_long(file_path)
    yc = [c for c in year_columns_from_dataframe(df)][:n_years]
    scen = df["Scenario"].astype(str).str.strip()

    def _row_values(variable_name: str) -> list:
        r = df[(df["Variable"] == variable_name) & (scen == label)]
        if r.empty:
            return [0.0] * n_years
        row = r.iloc[0]
        vals = []
        for c in yc:
            v = row.get(c, 0)
            vals.append(0.0 if pd.isna(v) else float(v))
        while len(vals) < n_years:
            vals.append(0.0)
        return vals[:n_years]

    opex_sub = df[(df["Variable"] == "OPEX") & (scen == label)]
    var_cost = []
    for c in yc:
        s = 0.0
        for _, row in opex_sub.iterrows():
            v = row.get(c, 0)
            s += 0.0 if pd.isna(v) else float(v)
        var_cost.append(s)
    while len(var_cost) < n_years:
        var_cost.append(0.0)
    var_cost = var_cost[:n_years]

    emissions = _row_values("Emissions")
    r_c = df[(df["Variable"] == "Carbon Price") & (scen == label)]
    if r_c.empty:
        carbon = [0.0] * n_years
    else:
        cr = r_c.iloc[0]
        raw_c = [np.nan if pd.isna(cr.get(c, np.nan)) else float(cr[c]) for c in yc]
        carbon = pd.Series(raw_c).ffill().bfill().fillna(0.0).tolist()[:n_years]
    while len(carbon) < n_years:
        carbon.append(0.0)
    carbon = carbon[:n_years]

    return {
        "variable_cost": var_cost,
        "fixed_cost": [0.0] * n_years,
        "co2_emission": emissions,
        "carbon_price": carbon,
        "carbon_credit_price": list(carbon),
    }


def emission_savings_series_from_investment_plan(
    file_path: str, totals_index: pd.Index
) -> pd.Series:
    df = read_investment_plan_long(file_path)
    v = df["Variable"].astype(str).str.strip()
    mask = v.eq("Emissions Savings") | v.str.endswith("Emissions Savings")
    sub = df.loc[mask]
    if sub.empty:
        return pd.Series(0.0, index=totals_index, dtype=float, name="emission_savings")
    label = pure_input_scenario_label("net_zero", file_path)
    nz = sub[sub["Scenario"].astype(str).str.strip().eq(label)]
    row = nz.iloc[0] if len(nz) else sub.iloc[0]
    vals_by_year: dict[int, float] = {}
    for c in year_columns_from_dataframe(df):
        y = int(float(c))
        val = row.get(c, 0)
        vals_by_year[y] = 0.0 if pd.isna(val) else float(val)
    out: list[float] = []
    for y in totals_index:
        yi = pd.to_numeric(y, errors="coerce")
        if pd.isna(yi):
            out.append(0.0)
        else:
            out.append(vals_by_year.get(int(yi), 0.0))
    return pd.Series(out, index=totals_index, dtype=float, name="emission_savings")
